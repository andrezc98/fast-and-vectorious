"""Flag text that overflows its box when rendered in DM Sans (the deck font), without needing PowerPoint.

LibreOffice on the build Mac falls back to a narrower serif, so renders hide overflow. This wraps every paragraph with
the real DM Sans glyph widths and compares the wrapped height with the box (or table row) height.

usage: python fit_check.py deck.pptx <dir with DMSans-Regular.ttf, DMSans-Bold.ttf, DMSansExtraBold-Regular.ttf>
"""
import os
import sys

from PIL import ImageFont
from pptx import Presentation
from pptx.util import Emu

EM_LINE = 1.302  # DM Sans (ascent + descent) / unitsPerEm: PowerPoint's single line height
_cache = {}


def face(fdir, name, bold, pt):
    f = "DMSansExtraBold-Regular.ttf" if "ExtraBold" in (name or "") else ("DMSans-Bold.ttf" if bold else "DMSans-Regular.ttf")
    key = (f, pt)
    if key not in _cache:
        _cache[key] = ImageFont.truetype(os.path.join(fdir, f), size=int(pt * 20))  # 20 px per pt: sub-point precision
    return _cache[key]


def para_height(fdir, p, width_in, default_pt=18):
    """Wrapped height of one paragraph in inches (greedy word wrap, like PowerPoint)."""
    words = []  # (text, font, pt)
    for r in p.runs:
        pt = r.font.size.pt if r.font.size else default_pt
        fnt = face(fdir, r.font.name, r.font.bold, pt)
        for i, w in enumerate(r.text.split(" ")):
            words.append(((" " if i else "") + w, fnt, pt))
    if not words:
        return default_pt * EM_LINE / 72, 0, 1
    lines, cur, maxpt, too_wide = 1, 0.0, 0, 0
    line_pts = []
    for t, fnt, pt in words:
        w = fnt.getlength(t) / 20 / 72  # inches
        if cur and cur + w > width_in:
            lines += 1
            line_pts.append(maxpt)
            cur, maxpt, w = 0.0, 0, fnt.getlength(t.lstrip()) / 20 / 72
        too_wide = max(too_wide, w - width_in)
        cur += w
        maxpt = max(maxpt, pt)
    line_pts.append(maxpt)
    space = p.space_after.pt if p.space_after is not None else 0
    return sum(pt * EM_LINE for pt in line_pts) / 72 + space / 72, too_wide, lines


def frame_height(fdir, tf, width_in):
    h, wide, wraps = 0.0, 0.0, False
    for p in tf.paragraphs:
        ph, w, n = para_height(fdir, p, width_in)
        h, wide, wraps = h + ph, max(wide, w), wraps or n > 1
    return h, wide, wraps


def check(path, fdir, tol=0.03):
    bad = []
    for i, s in enumerate(Presentation(path).slides, 1):
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                tf = sh.text_frame
                w = Emu(sh.width - tf.margin_left - tf.margin_right).inches
                hbox = Emu(sh.height - tf.margin_top - tf.margin_bottom).inches
                h, wide, wraps = frame_height(fdir, tf, w)
                if (wraps and h > hbox + tol) or wide > tol:  # a single unwrapped line just spills past a short box: harmless
                    bad.append((i, f"{tf.text[:60]!r}", f"needs {h:.2f} in, box {hbox:.2f}" + (f", a word is {wide:.2f} in too wide" if wide > tol else "")))
            if sh.has_table:
                tb = sh.table
                for r, row in enumerate(tb.rows):
                    for c, cell in enumerate(row.cells):
                        w = Emu(tb.columns[c].width - cell.margin_left - cell.margin_right).inches
                        hbox = Emu(row.height - cell.margin_top - cell.margin_bottom).inches
                        h, wide, wraps = frame_height(fdir, cell.text_frame, w)
                        if (wraps and h > hbox + tol) or wide > tol:
                            bad.append((i, f"table r{r}c{c} {cell.text[:40]!r}", f"needs {h:.2f} in, row {hbox:.2f}"))
    return bad


if __name__ == "__main__":
    out = check(sys.argv[1], sys.argv[2])
    for b in out:
        print(f"slide {b[0]:>2}  {b[1]}  {b[2]}")
    print(f"fit_check: {len(out)} overflow(s)")
